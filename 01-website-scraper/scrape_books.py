"""Scrape a whole online catalogue into a formatted, client-ready Excel workbook.

Target: https://books.toscrape.com, a public sandbox built for scraping practice.
Crawls every category (following pagination), collects title, price, rating,
stock and links for each product, then writes:

  * Products  - every product as a filterable, formatted table
  * Summary   - per-category counts, average price and rating, with a chart
  * Best deals - highly rated products under a price threshold

Usage:
    py -3.10 scrape_books.py                     # full catalogue -> books_catalogue.xlsx
    py -3.10 scrape_books.py --max-categories 3  # quick test run
"""

import argparse
import time
from dataclasses import dataclass
from datetime import datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
REQUEST_DELAY_S = 0.2  # be polite to the server

HEADER_FILL = PatternFill("solid", fgColor="1F3D2B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)


@dataclass
class Product:
    title: str
    category: str
    price: float
    rating: int
    in_stock: bool
    url: str


def fetch(session: requests.Session, url: str, retries: int = 3) -> BeautifulSoup:
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, timeout=20)
            resp.raise_for_status()
            resp.encoding = "utf-8"
            time.sleep(REQUEST_DELAY_S)
            return BeautifulSoup(resp.text, "html.parser")
        except requests.RequestException:
            if attempt == retries:
                raise
            time.sleep(2 * attempt)


def list_categories(session: requests.Session) -> list[tuple[str, str]]:
    soup = fetch(session, BASE_URL)
    links = soup.select("div.side_categories ul li ul li a")
    # The site files unsorted books under "Default"; give clients a clearer label
    rename = {"Default": "Uncategorised"}
    return [(rename.get(name := a.get_text(strip=True), name), urljoin(BASE_URL, a["href"])) for a in links]


def scrape_category(session: requests.Session, name: str, url: str) -> list[Product]:
    products = []
    while url:
        soup = fetch(session, url)
        for card in soup.select("article.product_pod"):
            link = card.select_one("h3 a")
            rating_word = card.select_one("p.star-rating")["class"][1]
            price_text = card.select_one("p.price_color").get_text(strip=True)
            products.append(Product(
                title=link["title"],
                category=name,
                price=float(price_text.lstrip("£Â")),
                rating=RATING_WORDS.get(rating_word, 0),
                in_stock="In stock" in card.select_one("p.availability").get_text(),
                url=urljoin(url, link["href"]),
            ))
        next_link = soup.select_one("li.next a")
        url = urljoin(url, next_link["href"]) if next_link else None
    return products


def style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_products_sheet(ws, products: list[Product]) -> None:
    headers = ["Title", "Category", "Price (£)", "Rating", "Stars", "In stock", "Link"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for p in products:
        ws.append([p.title, p.category, p.price, p.rating,
                   "★" * p.rating + "☆" * (5 - p.rating),
                   "Yes" if p.in_stock else "No", p.url])
        link_cell = ws.cell(row=ws.max_row, column=7)
        link_cell.hyperlink = p.url
        link_cell.value = "Open"
        link_cell.font = Font(color="0563C1", underline="single")

    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=3, max_col=3):
        row[0].number_format = "£#,##0.00"
    # Green = cheap, red = expensive, so outliers jump out at a glance
    ws.conditional_formatting.add(
        f"C2:C{last}",
        ColorScaleRule(start_type="min", start_color="63BE7B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="F8696B"),
    )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{last}"
    set_widths(ws, [60, 22, 12, 9, 10, 10, 9])


def write_summary_sheet(ws, products: list[Product], scraped_at: str) -> None:
    ws["A1"] = "Catalogue summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{len(products)} products scraped from {BASE_URL} on {scraped_at}"

    by_cat: dict[str, list[Product]] = {}
    for p in products:
        by_cat.setdefault(p.category, []).append(p)
    rows = sorted(
        ((cat, len(items),
          sum(i.price for i in items) / len(items),
          sum(i.rating for i in items) / len(items),
          min(i.price for i in items),
          max(i.price for i in items))
         for cat, items in by_cat.items()),
        key=lambda r: r[1], reverse=True,
    )

    header_row = 4
    headers = ["Category", "Products", "Avg price (£)", "Avg rating", "Cheapest (£)", "Priciest (£)"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=h)
    style_header(ws, header_row, len(headers))
    for r in rows:
        ws.append(list(r))
    last = ws.max_row
    for row in ws.iter_rows(min_row=header_row + 1, max_row=last, min_col=3, max_col=6):
        row[0].number_format = row[2].number_format = row[3].number_format = "£#,##0.00"
        row[1].number_format = "0.0"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    set_widths(ws, [24, 11, 14, 12, 14, 14])

    # Chart the 15 biggest categories
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Average price - 15 largest categories"
    chart.y_axis.title = "£"
    chart.legend = None
    top_last = min(last, header_row + 15)
    chart.add_data(Reference(ws, min_col=3, min_row=header_row, max_row=top_last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=top_last))
    chart.height, chart.width = 10, 18
    ws.add_chart(chart, "H4")


def write_deals_sheet(ws, products: list[Product], max_price: float) -> None:
    deals = sorted((p for p in products if p.rating >= 4 and p.price <= max_price and p.in_stock),
                   key=lambda p: (-p.rating, p.price))
    ws["A1"] = f"Best deals: 4-5 stars, in stock, under £{max_price:.0f}"
    ws["A1"].font = TITLE_FONT
    headers = ["Title", "Category", "Price (£)", "Stars", "Link"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=h)
    style_header(ws, 3, len(headers))
    for p in deals:
        ws.append([p.title, p.category, p.price, "★" * p.rating + "☆" * (5 - p.rating), "Open"])
        ws.cell(row=ws.max_row, column=3).number_format = "£#,##0.00"
        link_cell = ws.cell(row=ws.max_row, column=5)
        link_cell.hyperlink = p.url
        link_cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A4"
    set_widths(ws, [60, 22, 12, 10, 9])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", default="books_catalogue.xlsx")
    parser.add_argument("--max-categories", type=int, default=None, help="limit for quick test runs")
    parser.add_argument("--deal-price", type=float, default=20.0)
    args = parser.parse_args()

    session = requests.Session()
    session.headers["User-Agent"] = "Mozilla/5.0 (portfolio demo scraper)"

    started = time.time()
    categories = list_categories(session)[: args.max_categories]
    products: list[Product] = []
    for i, (name, url) in enumerate(categories, start=1):
        found = scrape_category(session, name, url)
        products.extend(found)
        print(f"[{i:>2}/{len(categories)}] {name:<22} {len(found):>4} products")

    scraped_at = datetime.now().strftime("%d %b %Y, %H:%M")
    wb = Workbook()
    write_summary_sheet(wb.active, products, scraped_at)
    wb.active.title = "Summary"
    write_products_sheet(wb.create_sheet("Products"), products)
    write_deals_sheet(wb.create_sheet("Best deals"), products, args.deal_price)
    wb.save(args.out)

    print(f"\nDone: {len(products)} products from {len(categories)} categories "
          f"in {time.time() - started:.0f}s -> {args.out}")


if __name__ == "__main__":
    main()

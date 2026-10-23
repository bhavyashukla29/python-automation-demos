"""Turn three messy sales exports into one clean, formatted Excel report.

Reads input/shopify_orders.csv, input/amazon_orders.csv and input/store_pos.csv,
fixes everything a person would otherwise fix by hand, and writes:

  * output/sales_report.xlsx - Summary (KPIs + charts), By product, Clean data,
                               and a Data quality log of every fix applied
  * output/clean_sales.csv   - the merged, clean dataset (feeds the dashboard demo)

Usage:
    py -3.10 build_report.py
"""

import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
INPUT, OUTPUT = HERE / "input", HERE / "output"

# Master price list: the single source of truth for names and prices
CATALOGUE = {
    "TEA-MC-250": ("Masala Chai Blend 250g", 349),
    "TEA-AG-500": ("Assam Gold Tea 500g", 449),
    "TEA-KK-100": ("Kashmiri Kahwa 100g", 399),
    "COF-FC-500": ("Filter Coffee Powder 500g", 599),
    "COF-CB-010": ("Cold Brew Coffee Bags (10)", 499),
    "EQP-FP-600": ("French Press 600ml", 1299),
    "EQP-CC-004": ("Ceramic Chai Cups (Set of 4)", 899),
    "GFT-TL-001": ("Tea Lovers Gift Hamper", 1799),
}
NAME_TO_SKU = {name: sku for sku, (name, _) in CATALOGUE.items()}

HEADER_FILL = PatternFill("solid", fgColor="1F3D2B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
KPI_FILL = PatternFill("solid", fgColor="EEF4F0")
THIN = Side(style="thin", color="C9D6CD")
RUPEES = '"₹"#,##0'

quality_log: list[tuple[str, str, int, str]] = []


def log(source: str, issue: str, count: int, action: str) -> None:
    if count:
        quality_log.append((source, issue, count, action))


# ---------- product name matching ----------

def tokens(text: str) -> set[str]:
    text = text.lower().replace("'s", "s")
    text = re.sub(r"(\d+)\s*(g|ml)\b", r"\1", text)  # "250g" / "250 g" -> "250"
    text = re.sub(r"\bx(\d+)\b", r"\1", text)         # "x10" -> "10"
    return set(re.findall(r"[a-z]+|\d+", text)) - {"of", "set", "the"}


CANONICAL_TOKENS = {name: tokens(name) for name in NAME_TO_SKU}


def match_product(raw: str) -> tuple[str | None, float]:
    """Best catalogue match for a free-text item name, with a 0-1 confidence."""
    raw_t = tokens(raw)
    if not raw_t:
        return None, 0.0
    best, best_score = None, 0.0
    for name, canon_t in CANONICAL_TOKENS.items():
        overlap = len(raw_t & canon_t)
        # how much of the raw text is explained, tie-broken by how much of the name is covered
        score = overlap / len(raw_t) + 0.01 * overlap / len(canon_t)
        if score > best_score:
            best, best_score = name, score
    return (best, min(best_score, 1.0)) if best_score >= 0.5 else (None, best_score)


# ---------- loaders: one per channel, each returns the common schema ----------

def load_shopify() -> pd.DataFrame:
    df = pd.read_csv(INPUT / "shopify_orders.csv", encoding="utf-8-sig")
    src = "Shopify"
    dupes = df.duplicated().sum()
    df = df.drop_duplicates()
    log(src, "Duplicate rows from overlapping exports", dupes, "Removed")
    refunded = (df["Financial Status"] == "refunded").sum()
    df = df[df["Financial Status"] == "paid"]
    log(src, "Refunded orders", refunded, "Excluded from revenue")
    return pd.DataFrame({
        "date": pd.to_datetime(df["Created at"].str[:10], format="%Y-%m-%d"),
        "channel": "Website",
        "order_id": df["Name"],
        "sku": df["Lineitem sku"],
        "qty": df["Lineitem quantity"].astype(int),
        "revenue": df["Lineitem price"].astype(float) * df["Lineitem quantity"].astype(int),
    })


def load_amazon() -> pd.DataFrame:
    df = pd.read_csv(INPUT / "amazon_orders.csv", encoding="utf-8-sig")
    src = "Amazon"
    cancelled = (df["order-status"] == "Cancelled").sum()
    df = df[df["order-status"] != "Cancelled"]
    log(src, "Cancelled orders", cancelled, "Excluded from revenue")
    canonical = df["sku"].map(lambda s: CATALOGUE[s][0])
    renamed = (df["product-name"] != canonical).sum()
    log(src, "Misspelt / inconsistent product names", renamed, "Standardised using SKU")
    log(src, 'Prices stored as text ("₹1,299.00")', len(df), "Converted to numbers")
    return pd.DataFrame({
        "date": pd.to_datetime(df["purchase-date"], format="%d/%m/%Y"),
        "channel": "Amazon",
        "order_id": df["order-id"],
        "sku": df["sku"],
        "qty": df["quantity-purchased"].astype(int),
        "revenue": df["item-price"].str.replace(r"[₹,]", "", regex=True).astype(float),
    })


def parse_mixed_date(text: str) -> pd.Timestamp:
    for fmt in ("%d %b %Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return pd.Timestamp(datetime.strptime(text.strip(), fmt))
        except ValueError:
            continue
    return pd.NaT


def load_pos() -> pd.DataFrame:
    df = pd.read_csv(INPUT / "store_pos.csv", encoding="utf-8-sig", dtype=str, keep_default_na=False)
    src = "Store POS"
    blank = (df.apply(lambda r: "".join(r).strip(), axis=1) == "").sum()
    df = df[df.apply(lambda r: "".join(r).strip(), axis=1) != ""].copy()
    log(src, "Blank rows", blank, "Removed")

    df["date"] = df["Date"].map(parse_mixed_date)
    formats = df["Date"].str.contains("/").sum() + df["Date"].str.contains("-").sum()
    log(src, "Dates in 3 different formats", int(formats), "Unified to one date format")
    bad_dates = df["date"].isna().sum()
    df = df[df["date"].notna()]
    log(src, "Unreadable dates", bad_dates, "Removed")

    matches = df["Item"].map(match_product)
    df["product"] = matches.map(lambda m: m[0])
    unmatched = df["product"].isna().sum()
    log(src, "Abbreviated / lowercase item names", int((df["Item"].str.strip() != df["product"]).sum()),
        "Matched to the price list")
    log(src, "Item names that could not be matched", unmatched, "Removed - needs manual review")
    df = df[df["product"].notna()]

    df["qty"] = df["Qty"].astype(int)
    df["sku"] = df["product"].map(NAME_TO_SKU)
    list_price = df["sku"].map(lambda s: CATALOGUE[s][1]) * df["qty"]
    missing = (df["Amount"].str.strip() == "").sum()
    df["revenue"] = pd.to_numeric(df["Amount"], errors="coerce").fillna(list_price)
    log(src, "Missing sale amounts", missing, "Recovered from price list x quantity")

    order_ids = "POS-" + df["date"].dt.strftime("%Y%m%d") + "-" + (df.groupby("date").cumcount() + 1).astype(str)
    return pd.DataFrame({"date": df["date"], "channel": "Store", "order_id": order_ids,
                         "sku": df["sku"], "qty": df["qty"], "revenue": df["revenue"]})


# ---------- Excel output ----------

def header(ws, row: int, values: list[str], col: int = 1) -> None:
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, values: list[int]) -> None:
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_summary(ws, sales: pd.DataFrame) -> None:
    period = f"{sales['date'].min():%d %b %Y} - {sales['date'].max():%d %b %Y}"
    ws["A1"] = "Sales report - all channels"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{period}  ·  generated automatically on {datetime.now():%d %b %Y, %H:%M}"
    ws["A2"].font = Font(italic=True, color="666666")

    orders = sales["order_id"].nunique()
    kpis = [("Revenue", sales["revenue"].sum(), RUPEES), ("Orders", orders, "#,##0"),
            ("Units sold", sales["qty"].sum(), "#,##0"), ("Avg order value", sales["revenue"].sum() / orders, RUPEES)]
    for i, (label, value, fmt) in enumerate(kpis):
        col = 1 + i * 2
        lab = ws.cell(row=4, column=col, value=label)
        val = ws.cell(row=5, column=col, value=float(value))
        lab.font = Font(color="666666")
        val.font, val.number_format = Font(bold=True, size=14), fmt
        for c in (lab, val):
            c.fill = KPI_FILL
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

    # Monthly revenue by channel
    monthly = (sales.assign(month=sales["date"].dt.to_period("M"))
               .pivot_table(index="month", columns="channel", values="revenue", aggfunc="sum", fill_value=0))
    channels = list(monthly.columns)
    top = 8
    header(ws, top, ["Month"] + channels + ["Total"])
    for r, (month, row) in enumerate(monthly.iterrows(), start=top + 1):
        ws.cell(row=r, column=1, value=month.strftime("%b %Y"))
        for c, ch in enumerate(channels, start=2):
            ws.cell(row=r, column=c, value=float(row[ch])).number_format = RUPEES
        ws.cell(row=r, column=len(channels) + 2, value=float(row.sum())).number_format = RUPEES
    last = top + len(monthly)

    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title = "Monthly revenue by channel"
    chart.add_data(Reference(ws, min_col=2, max_col=len(channels) + 1, min_row=top, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "H4")

    # Top products
    prod_top = last + 3
    by_product = (sales.groupby("product").agg(units=("qty", "sum"), revenue=("revenue", "sum"))
                  .sort_values("revenue", ascending=False))
    header(ws, prod_top, ["Product", "Units", "Revenue", "Share"])
    total = by_product["revenue"].sum()
    for r, (name, row) in enumerate(by_product.iterrows(), start=prod_top + 1):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=int(row["units"])).number_format = "#,##0"
        ws.cell(row=r, column=3, value=float(row["revenue"])).number_format = RUPEES
        ws.cell(row=r, column=4, value=float(row["revenue"] / total)).number_format = "0.0%"
    prod_last = prod_top + len(by_product)

    bar = BarChart()
    bar.type, bar.title, bar.legend = "bar", "Revenue by product", None
    bar.add_data(Reference(ws, min_col=3, min_row=prod_top, max_row=prod_last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=prod_top + 1, max_row=prod_last))
    bar.y_axis.scaling.orientation = "minMax"
    bar.x_axis.scaling.orientation = "maxMin"  # biggest at the top
    bar.height, bar.width = 8, 16
    ws.add_chart(bar, "H21")
    widths(ws, [30, 14, 14, 14, 14, 14, 14])


def write_by_product(ws, sales: pd.DataFrame) -> None:
    pivot = (sales.assign(month=sales["date"].dt.to_period("M"))
             .pivot_table(index="product", columns="month", values="revenue", aggfunc="sum", fill_value=0))
    pivot["Total"] = pivot.sum(axis=1)
    pivot = pivot.sort_values("Total", ascending=False)
    header(ws, 1, ["Product"] + [c.strftime("%b %Y") if c != "Total" else c for c in pivot.columns])
    for r, (name, row) in enumerate(pivot.iterrows(), start=2):
        ws.cell(row=r, column=1, value=name)
        for c, v in enumerate(row, start=2):
            ws.cell(row=r, column=c, value=float(v)).number_format = RUPEES
    ws.freeze_panes = "B2"
    widths(ws, [30] + [13] * len(pivot.columns))


def write_clean_data(ws, sales: pd.DataFrame) -> None:
    cols = ["date", "channel", "order_id", "product", "sku", "qty", "revenue"]
    header(ws, 1, ["Date", "Channel", "Order ID", "Product", "SKU", "Qty", "Revenue"])
    for r in sales[cols].itertuples(index=False):
        ws.append([r.date.to_pydatetime(), r.channel, r.order_id, r.product, r.sku, int(r.qty), float(r.revenue)])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].number_format = "DD-MMM-YYYY"
        row[6].number_format = RUPEES
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    widths(ws, [13, 10, 22, 30, 13, 6, 12])


def write_quality_log(ws) -> None:
    ws["A1"] = "Data quality log - every fix applied automatically"
    ws["A1"].font = Font(bold=True, size=14)
    header(ws, 3, ["Source", "Issue found", "Rows", "Action taken"])
    for entry in quality_log:
        ws.append(list(entry))
    widths(ws, [12, 42, 8, 38])


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    sales = pd.concat([load_shopify(), load_amazon(), load_pos()], ignore_index=True)
    sales["product"] = sales["sku"].map(lambda s: CATALOGUE[s][0])
    sales = sales.sort_values(["date", "channel"]).reset_index(drop=True)

    wb = Workbook()
    wb.active.title = "Summary"
    write_summary(wb.active, sales)
    write_by_product(wb.create_sheet("By product"), sales)
    write_clean_data(wb.create_sheet("Clean data"), sales)
    write_quality_log(wb.create_sheet("Data quality log"))
    wb.save(OUTPUT / "sales_report.xlsx")
    sales.to_csv(OUTPUT / "clean_sales.csv", index=False)

    print(f"{len(sales):,} clean sale lines · ₹{sales['revenue'].sum():,.0f} revenue")
    print("Fixes applied:")
    for source, issue, count, action in quality_log:
        print(f"  [{source}] {issue}: {count} -> {action}")
    print(f"-> {OUTPUT / 'sales_report.xlsx'}")


if __name__ == "__main__":
    main()
